"""
Build the synthetic enterprise for the Meridian demo from fixtures/scenario.py.

    python fixtures/build_scenario.py            # build everything that is missing
    python fixtures/build_scenario.py --force    # rebuild the git repositories too

Outputs (all synthetic):
  .meridian/repos/<component>/           git repository per component, tagged per version
  environments/<env>/adapters/*.json     raw read-only adapter outputs (kubectl, Liberty, Flyway, ...)
  environments/deployment-events.json    pipeline / GitOps / CAB deployment history
  environments/validation-runs.json      validation evidence from test pipelines
  environments/scheduled-changes.json    approved but not yet executed changes
  documents/*.pdf|docx|xlsx               change request, release notes, interface control document
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import shutil
import subprocess
import sys
from datetime import datetime, timezone
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from fixtures import scenario as S  # noqa: E402

SAMPLE = ROOT / "sample-system"
MERIDIAN = ROOT / ".meridian"
REPOS = MERIDIAN / "repos"
ENVS = ROOT / "environments"
DOCS = ROOT / "documents"

COMPONENTS = {c["id"]: c for c in S.COMPONENTS}


def _iso(ts: str) -> datetime:
    return datetime.fromisoformat(ts.replace("Z", "+00:00"))


def _z(dt: datetime) -> str:
    return dt.astimezone(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")


def _write_json(path: Path, data) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, indent=2) + "\n", encoding="utf-8")


def tag_for(version: str) -> str:
    return version if version[:1] in ("v", "V") else f"v{version}"


# ============================================================================ git repos

def _git(repo: Path, *args: str, env: dict | None = None) -> str:
    hooks = MERIDIAN / ".nohooks"
    hooks.mkdir(parents=True, exist_ok=True)
    cmd = [
        "git", "-c", "core.autocrlf=false", "-c", "core.safecrlf=false",
        "-c", "commit.gpgsign=false", "-c", "tag.gpgsign=false",
        "-c", f"core.hooksPath={hooks.as_posix()}", *args,
    ]
    result = subprocess.run(cmd, cwd=repo, capture_output=True, text=True, env=env)
    if result.returncode != 0:
        raise RuntimeError(f"git {' '.join(args)} failed: {result.stderr}")
    return result.stdout.strip()


def _repo_readme(component: str) -> str:
    c = COMPONENTS[component]
    return (
        f"# {component}\n\n"
        f"Owner: {c['team']}  \n"
        f"Runtime: {c['runtime']} on {c['platform']} ({c['cloud']})  \n"
        f"Deployed by: {c['deploy_tool']}\n\n"
        "Synthetic repository generated for the Meridian demo. All content is fictional.\n"
    )


def build_repos(force: bool = False) -> dict[str, dict[str, str]]:
    """Create one git repository per component. Returns {component: {version: commit}}."""
    REPOS.mkdir(parents=True, exist_ok=True)
    commits: dict[str, dict[str, str]] = {}
    for component, spec in S.GIT_HISTORY.items():
        repo = REPOS / component
        expected_tags = [tag_for(v) for v, _ in spec["versions"]]
        if repo.exists() and not force:
            try:
                commits[component] = {
                    v: _git(repo, "rev-parse", "--short=7", f"{tag_for(v)}^{{commit}}")
                    for v, _ in spec["versions"]
                }
                continue
            except RuntimeError:
                pass  # incomplete repo, rebuild below
        if repo.exists():
            shutil.rmtree(repo)
        repo.mkdir(parents=True)
        _git(repo, "init", "-q", "-b", "main")

        name, email = S.PEOPLE[spec["owner"]]
        first = True
        commits[component] = {}
        for version, version_commits in spec["versions"]:
            vdir = SAMPLE / component / "versions" / version
            vfiles = {
                p.relative_to(vdir).as_posix(): p.read_text(encoding="utf-8").replace("\r\n", "\n")
                for p in sorted(vdir.rglob("*"))
                if p.is_file() and "__pycache__" not in p.parts and p.suffix != ".pyc"
            }
            for date, message, files in version_commits:
                chosen = vfiles if files is None else {k: vfiles[k] for k in files}
                if first:
                    chosen = {"README.md": _repo_readme(component), **chosen}
                    first = False
                for rel, content in chosen.items():
                    target = repo / rel
                    target.parent.mkdir(parents=True, exist_ok=True)
                    with open(target, "w", encoding="utf-8", newline="\n") as fh:
                        fh.write(content)
                env = {
                    **os.environ,
                    "GIT_AUTHOR_NAME": name, "GIT_AUTHOR_EMAIL": email, "GIT_AUTHOR_DATE": date,
                    "GIT_COMMITTER_NAME": name, "GIT_COMMITTER_EMAIL": email, "GIT_COMMITTER_DATE": date,
                }
                _git(repo, "add", "-A", env=env)
                _git(repo, "commit", "-q", "--allow-empty", "-m", message, env=env)
            _git(repo, "tag", tag_for(version))
            commits[component][version] = _git(repo, "rev-parse", "--short=7", "HEAD")
        assert all(t for t in expected_tags)
    return commits


# ==================================================================== environment state

def _baseline_time(env: str, component: str) -> str:
    if env == "prod":
        return S.PROD_BASELINE_DEPLOYED_AT[component]
    return S.BASELINE_DEPLOYED_AT[env]


def composition_at(env: str, at: str) -> dict[str, tuple[str, str]]:
    """{component: (version, deployed_at)} for an environment at a point in time."""
    state = {c: (S.RELEASE["baseline"][c], _baseline_time(env, c)) for c in COMPONENTS}
    for e_env, comp, _frm, to, ts, _note in sorted(S.DEPLOYMENTS, key=lambda d: d[4]):
        if e_env == env and _iso(ts) <= _iso(at):
            state[comp] = (to, ts)
    return state


def _digest(*parts: str) -> str:
    return "sha256:" + hashlib.sha256("|".join(parts).encode()).hexdigest()


ENV_PROFILE = {
    "dev":   {"replicas": 1, "cpu": "250m", "mem": "512Mi", "log": "DEBUG", "pool": "general-dev", "hpa": (1, 2)},
    "test":  {"replicas": 2, "cpu": "250m", "mem": "512Mi", "log": "DEBUG", "pool": "general-test", "hpa": (2, 3)},
    "stage": {"replicas": 3, "cpu": "500m", "mem": "1Gi", "log": "INFO", "pool": "general-stage", "hpa": (3, 6)},
    "prod":  {"replicas": 6, "cpu": "1000m", "mem": "2Gi", "log": "WARN", "pool": "general-prod", "hpa": (6, 24)},
}

SERVICE_ENV = {
    "web-checkout": {"ORDER_API_URL": "https://order-api.{env}.acme.internal/v2/orders", "SESSION_TTL_SECONDS": {"prod": "1800"}},
    "order-api": {"PAYMENT_GRPC_TARGET": "payment-service.{env}.acme.internal:443", "DB_POOL_SIZE": {"prod": "40", "stage": "20"}},
    "payment-service": {"KAFKA_BOOTSTRAP": "lkc-{env}-orders.confluent.cloud:9092", "PSP_TIMEOUT_MS": {"prod": "2500", "stage": "4000"}},
    "billing-service": {"MQ_QMGR": "QM.INTEG.01", "MQ_CONN": "mq-{env}-dc2-01(1414)", "BATCH_SIZE": {"prod": "200", "stage": "50"}},
    "mq-bridge": {"MQ_QMGR": "QM.INTEG.01", "MQ_CONN": "mq-{env}-dc2-01(1414)", "LEDGER_RECORD_LAYOUT": None},
}

FEATURE_FLAGS = {
    "web-checkout": {"prod": "express-pay=on,gift-cards=on", "stage": "express-pay=off,gift-cards=on"},
    "order-api": {"prod": "fraud-precheck=on", "stage": "fraud-precheck=on"},
}


def _k8s_deployment(env: str, name: str, version: str, commit: str, deployed_at: str,
                    registry: str, namespace: str, flow_component: bool, replicas_override: int | None = None) -> dict:
    prof = ENV_PROFILE[env]
    replicas = replicas_override or prof["replicas"]
    env_vars = [{"name": "LOG_LEVEL", "value": prof["log"]},
                {"name": "OTEL_EXPORTER_OTLP_ENDPOINT", "value": f"http://otel-collector.{env}.acme.internal:4317"}]
    for key, value in SERVICE_ENV.get(name, {}).items():
        if value is None:
            value = "v7" if version == "3.1" else "v6"
        elif isinstance(value, dict):
            value = value.get(env, value.get("stage", "10"))
        env_vars.append({"name": key, "value": value.format(env=env)})
    flags = FEATURE_FLAGS.get(name, {}).get(env if env in ("prod", "stage") else "stage")
    if flags:
        env_vars.append({"name": "FEATURE_FLAGS", "value": flags})
    image = f"{registry}/{name}:{version}@{_digest(name, version, commit)}"
    labels = {
        "app.kubernetes.io/name": name,
        "app.kubernetes.io/version": version,
        "app.kubernetes.io/managed-by": "argocd" if "argo" in COMPONENTS.get(name, {}).get("deploy_tool", "Argo").lower() else "helm",
    }
    if flow_component:
        labels["app.kubernetes.io/part-of"] = S.FLOW_ID
    return {
        "apiVersion": "apps/v1",
        "kind": "Deployment",
        "metadata": {
            "name": name,
            "namespace": namespace,
            "labels": labels,
            "annotations": {
                "deployment.kubernetes.io/revision": str(40 + int(hashlib.md5((name + version).encode()).hexdigest()[:2], 16) % 30),
                "org.opencontainers.image.revision": commit,
                "org.opencontainers.image.version": version,
            },
        },
        "spec": {
            "replicas": replicas,
            "template": {"spec": {
                "nodeSelector": {"acme.io/node-pool": prof["pool"]},
                "containers": [{
                    "name": name,
                    "image": image,
                    "resources": {
                        "requests": {"cpu": prof["cpu"], "memory": prof["mem"]},
                        "limits": {"cpu": prof["cpu"], "memory": prof["mem"]},
                    },
                    "env": env_vars,
                }],
            }},
        },
        "status": {
            "replicas": replicas, "readyReplicas": replicas, "availableReplicas": replicas, "updatedReplicas": replicas,
            "conditions": [
                {"type": "Available", "status": "True", "reason": "MinimumReplicasAvailable", "lastUpdateTime": deployed_at},
                {"type": "Progressing", "status": "True", "reason": "NewReplicaSetAvailable", "lastUpdateTime": deployed_at},
            ],
        },
    }


def _neighbour_version(kind: str, name: str, env: str) -> str:
    for n, versions in S.NEIGHBOURS.get(kind, []):
        if n == name:
            return versions.get(env, versions["stage"])
    raise KeyError(name)


def build_adapters(commits: dict[str, dict[str, str]]) -> None:
    for env in S.ENVIRONMENTS:
        state = composition_at(env, S.SNAPSHOT_TIME)
        out = ENVS / env / "adapters"
        out.mkdir(parents=True, exist_ok=True)

        # ---- Kubernetes clusters (EKS, AKS, GKE, OpenShift)
        clusters: dict[str, dict] = {}
        for comp_id, comp in COMPONENTS.items():
            if comp["adapter"] != "kubernetes":
                continue
            cluster = comp["cluster"].format(env=env)
            version, deployed_at = state[comp_id]
            entry = clusters.setdefault(cluster, {"platform": comp["platform"], "cloud": comp["cloud"],
                                                   "region": comp["location"], "items": []})
            prod_replicas = {"web-checkout": 6, "order-api": 6, "payment-service": 4, "billing-service": 4, "mq-bridge": 2}
            entry["items"].append(_k8s_deployment(
                env, comp_id, version, commits[comp_id][version], deployed_at,
                comp["registry"], comp["namespace"], True,
                prod_replicas[comp_id] if env == "prod" else None,
            ))
        kind_of = {"GKE": "gke", "EKS": "eks", "AKS": "aks", "OpenShift": "ocp"}
        for cluster, entry in clusters.items():
            kind = kind_of[entry["platform"]]
            for name, _ in S.NEIGHBOURS.get(kind, []):
                version = _neighbour_version(kind, name, env)
                entry["items"].append(_k8s_deployment(
                    env, name, version, hashlib.sha1(f"{name}{version}".encode()).hexdigest()[:7],
                    "2026-09-18T12:00:00Z" if env == "prod" else "2026-09-19T12:00:00Z",
                    "registry.acme.internal/shared", "shared", False,
                ))
            _write_json(out / f"kubernetes-{cluster}.json", {
                "adapter": "kubernetes",
                "source": f"kubectl --context {cluster} get deployments -A -o json",
                "collected_at": S.SNAPSHOT_TIME,
                "cluster": cluster, "platform": entry["platform"], "cloud": entry["cloud"], "region": entry["region"],
                "_synthetic": True,
                "apiVersion": "v1", "kind": "List", "items": entry["items"],
            })

        # ---- Confluent Schema Registry
        k_version, k_at = state["kafka-orders"]
        _write_json(out / "schema-registry.json", {
            "adapter": "schema-registry",
            "source": f"GET https://psrc-{env}-orders.acme.internal/subjects/orders-value/versions/latest",
            "collected_at": S.SNAPSHOT_TIME, "_synthetic": True,
            "cluster": f"lkc-{env}-orders", "platform": "Kafka", "cloud": "Confluent Cloud",
            "subjects": [{
                "subject": "orders-value", "topic": "orders.v1",
                "version": int(k_version[1:]), "id": 100200 + int(k_version[1:]),
                "compatibility": "BACKWARD",
                "registered_at": k_at,
                "schema_git_commit": commits["kafka-orders"][k_version],
            }],
            "config": {"retention.ms": "604800000" if env == "prod" else "259200000",
                       "partitions": 24 if env == "prod" else 6},
        })

        # ---- WebSphere Liberty inventory (VMware VM)
        l_version, l_at = state["legacy-ledger"]
        host = f"was-{env}-ldg-03.dc2.acme.internal"
        _write_json(out / "websphere-liberty.json", {
            "adapter": "websphere-liberty",
            "source": "Liberty Admin Center REST export + vSphere inventory",
            "collected_at": S.SNAPSHOT_TIME, "_synthetic": True,
            "servers": [{
                "host": host,
                "vm": {"vcenter": "vcsa-dc2.acme.internal", "cluster": "vsphere-dc2-ledger",
                       "cpu": 8 if env == "prod" else 4, "memory_gb": 32 if env == "prod" else 16},
                "liberty": {"version": "24.0.0.9", "java": "IBM Semeru Runtime 17.0.12"},
                "jvm": {"heap_max": "4096m" if env == "prod" else "2048m"},
                "applications": [{
                    "name": "ldgpost",
                    "archive": f"ldgpost-{l_version}.ear",
                    "state": "STARTED",
                    "deployed_at": l_at,
                    "manifest": {
                        "Implementation-Title": "LDGPOST",
                        "Implementation-Version": l_version,
                        "Implementation-Build": commits["legacy-ledger"][l_version],
                        "Implementation-Vendor": "Acme Core Ledger",
                    },
                    "mq": {"queue_manager": "QM.LEDG.01", "queue": "LEDGER.IN"},
                }],
            }],
        })

        # ---- Flyway history (Db2)
        db_version, db_at = state["ledger-db"]
        rows = [
            {"installed_rank": 13, "version": "13", "description": "ledger indexes", "script": "V13__ledger_indexes.sql",
             "installed_on": "2026-08-30T02:14:09Z", "success": True},
            {"installed_rank": 14, "version": "14", "description": "ledger transactions", "script": "V14__ledger_transactions.sql",
             "installed_on": _baseline_time(env, "ledger-db") if env != "prod" else "2026-09-21T18:10:00Z", "success": True},
        ]
        if db_version == "V15":
            rows.append({"installed_rank": 15, "version": "15", "description": "widen customer add currency",
                         "script": "V15__widen_customer_add_currency.sql", "installed_on": db_at, "success": True})
        for row in rows:
            row.update({"type": "SQL", "installed_by": "flyway_svc", "execution_time": 380 + row["installed_rank"] * 7,
                        "checksum": int(hashlib.md5(row["script"].encode()).hexdigest()[:7], 16)})
        _write_json(out / "flyway-ledgdb.json", {
            "adapter": "flyway",
            "source": f"SELECT * FROM LEDGER.\"flyway_schema_history\" @ db2-{env}-ldg01/LEDGDB",
            "collected_at": S.SNAPSHOT_TIME, "_synthetic": True,
            "database": "LEDGDB", "host": f"db2-{env}-ldg01.dc2.acme.internal", "schema": "LEDGER",
            "migrations_git_commit": commits["ledger-db"][db_version],
            "flyway_schema_history": rows,
        })

        # ---- IBM MQ queue manager configuration (not versioned; appears in naive diffs)
        _write_json(out / "ibm-mq.json", {
            "adapter": "ibm-mq",
            "source": "MQ REST API · GET /ibmmq/rest/v2/admin/qmgr/*/queue",
            "collected_at": S.SNAPSHOT_TIME, "_synthetic": True,
            "queue_managers": [
                {"name": "QM.INTEG.01", "host": f"mq-{env}-dc2-01", "queues": [
                    {"name": "BILL.OUT", "maxdepth": 50000 if env == "prod" else 5000, "defpersistence": "YES"},
                    {"name": "LEDGER.IN", "maxdepth": 50000 if env == "prod" else 5000, "defpersistence": "YES"},
                    {"name": "LEDGER.HOLD", "maxdepth": 20000 if env == "prod" else 2000, "defpersistence": "YES"},
                ]},
                {"name": "QM.LEDG.01", "host": f"mq-{env}-dc2-02", "queues": [
                    {"name": "LEDGER.IN", "maxdepth": 50000 if env == "prod" else 5000, "defpersistence": "YES",
                     "backout_threshold": 3 if env == "prod" else 5},
                ]},
            ],
        })


def build_history(commits: dict[str, dict[str, str]]) -> None:
    events = []
    counters: dict[str, int] = {}
    for env, comp, frm, to, ts, note in sorted(S.DEPLOYMENTS, key=lambda d: d[4]):
        tool, actor, kind = S.DEPLOY_TOOLING[comp]
        counters[comp] = counters.get(comp, 0) + 1
        run_no = 1800 + counters[comp] * 7 + len(comp)
        event = {
            "event_id": f"dep-{env}-{comp}-{to}".replace(".", "_"),
            "environment": env,
            "component": comp,
            "from_version": frm,
            "to_version": to,
            "commit": commits.get(comp, {}).get(to),
            "deployed_at": ts,
            "tool": tool,
            "trigger": kind,
            "actor": actor,
            "run": f"{kind}-{env}-{run_no}",
            "status": "SUCCEEDED",
            "note": note,
        }
        events.append(event)
    _write_json(ENVS / "deployment-events.json", {
        "_synthetic": True,
        "description": "Deployment history collected from Google Cloud Deploy, Argo CD, Azure DevOps, GitHub Actions and Jenkins/CAB records.",
        "baseline": {
            env: {c: {"version": S.RELEASE["baseline"][c], "deployed_at": _baseline_time(env, c),
                      "commit": commits[c].get(S.RELEASE["baseline"][c])}
                  for c in COMPONENTS}
            for env in S.ENVIRONMENTS
        },
        "events": events,
    })

    runs = []
    for run in S.VALIDATION_RUNS:
        edges = []
        for edge_id, (pv, cv) in run["pairs"].items():
            edge = next(e for e in S.EDGES if e["id"] == edge_id)
            semantic = S.SEMANTIC_ASSERTIONS[edge_id] if run["semantic"] else 0
            edges.append({
                "edge_id": edge_id,
                "producer": edge["producer"], "producer_version": pv,
                "producer_commit": commits[edge["producer"]].get(pv),
                "consumer": edge["consumer"], "consumer_version": cv,
                "consumer_commit": commits[edge["consumer"]].get(cv),
                "messages": run["messages"],
                "semantic_assertions": semantic,
                "semantic_assertions_passed": semantic if run["result"] == "PASS" else 0,
                "technical_checks": run.get("checks", ["transport ACK", "HTTP / consumer status"]),
            })
        runs.append({k: v for k, v in run.items() if k not in ("pairs", "semantic", "checks")} | {"edges": edges})
    _write_json(ENVS / "validation-runs.json", {"_synthetic": True, "runs": runs,
                                                "fixtures": S.LEDGER_FIXTURES})

    _write_json(ENVS / "scheduled-changes.json", {"_synthetic": True, "changes": [
        {**c, "changes": [{"component": a, "from_version": b, "to_version": d} for a, b, d in c["changes"]]}
        for c in S.SCHEDULED_CHANGES
    ]})

    _write_json(ENVS / "flows.json", {
        "_synthetic": True,
        "organization": S.ORG,
        "snapshot_time": S.SNAPSHOT_TIME,
        "environments": S.ENVIRONMENTS,
        "promotion_path": S.PROMOTION_PATH,
        "flows": [{
            "id": S.FLOW_ID, "name": S.FLOW_NAME, "description": S.FLOW_DESCRIPTION,
            "components": S.COMPONENTS, "edges": S.EDGES,
        }],
    })
    _write_json(ENVS / "releases.json", {"_synthetic": True, "releases": [S.RELEASE]})


# ============================================================================ documents

def build_documents(commits: dict[str, dict[str, str]]) -> None:
    DOCS.mkdir(exist_ok=True)
    cr = {
        "id": "CR-4471",
        "title": "Upgrade legacy-ledger to 7.0 and apply Db2 migration V15 (LEDGREC rev 7)",
        "release": "R-26.9",
        "requested_by": "Ruth Adeyemi (Core Ledger)",
        "change_type": "Normal",
        "risk": "Medium",
        "cab": "CAB-0924 · approved 2026-09-24",
        "window": "Sunday 2026-09-27 02:00-04:00 UTC",
        "status": "Scheduled",
        "components": [
            {"component": "ledger-db", "from": "V14", "to": "V15", "target": "db2-prod-ldg01 / LEDGDB"},
            {"component": "legacy-ledger", "from": "6.9", "to": "7.0", "target": "was-prod-ldg-03 (Liberty)"},
        ],
        "dependency": "mq-bridge 3.1 emits LEDGREC rev 7. This change MUST be completed before mq-bridge 3.1 is promoted to Production.",
        "implementation": [
            "Quiesce LEDGER.IN consumers (stop LDGPOST listener).",
            "Apply Flyway migration V15 to LEDGDB (widen CUSTOMER_ID, add CURRENCY_CODE).",
            "Deploy ldgpost-7.0.ear to was-prod-ldg-03 and restart Liberty.",
            "Resume LEDGER.IN consumers and verify postings for rev 6 and rev 7 records.",
        ],
        "backout": "Redeploy ldgpost-6.9.ear. V15 is additive; no schema rollback required.",
        "verification": "Post one rev 6 and one rev 7 test record; compare LEDGER.TRANSACTIONS rows with source events.",
    }
    notes = {
        "release": "R-26.9",
        "title": S.RELEASE["title"],
        "summary": S.RELEASE["requirement"],
        "changed_services": [
            {"component": c, "from": S.RELEASE["baseline"][c], "to": v, "team": COMPONENTS[c]["team"],
             "deployed_by": COMPONENTS[c]["deploy_tool"]}
            for c, v in S.RELEASE["target"].items() if S.RELEASE["baseline"][c] != v
        ],
        "deployment_dependencies": S.RELEASE["dependencies"],
        "rollout_order": S.RELEASE["rollout_order"],
        "migration_requirements": [
            "Flyway V15 on LEDGDB before legacy-ledger 7.0 starts.",
            "orders-value schema v4 registered before payment-service 8.4 deploys.",
        ],
        "compatibility_notes": [
            "LEDGREC rev 7 is 25 bytes: customerId(12) currencyCode(3) amount(10).",
            "legacy-ledger 7.0 selects the layout by record length and accepts rev 6 (23 bytes) and rev 7 (25 bytes).",
            "legacy-ledger 6.9 only knows rev 6. It is not supported with mq-bridge 3.1.",
        ],
        "validated_in": "Stage · stage-e2e-order-to-ledger · 2026-09-24 14:30 UTC · PASS",
    }
    icd = {
        "id": "LEDG-ICD-007",
        "title": "Interface Control Document: LEDGER.IN posting record (LEDGREC)",
        "owner": "Core Ledger",
        "rev7": [
            {"Field": "customerId", "Offset": 1, "Length": 12, "Type": "STRING", "Required": "YES", "InterfaceVersion": "rev 7",
             "Description": "Customer identifier, left aligned, space padded. Widened from 10 in R-26.9."},
            {"Field": "currencyCode", "Offset": 13, "Length": 3, "Type": "STRING", "Required": "YES", "InterfaceVersion": "rev 7",
             "Description": "ISO 4217 currency code. New in R-26.9."},
            {"Field": "amount", "Offset": 16, "Length": 10, "Type": "DECIMAL", "Required": "YES", "InterfaceVersion": "rev 7",
             "Description": "Amount, zero padded, explicit decimal point, 2 decimals."},
        ],
        "rev6": [
            {"Field": "customerId", "Offset": 1, "Length": 10, "Type": "STRING", "Required": "YES", "InterfaceVersion": "rev 6",
             "Description": "Customer identifier, left aligned, space padded."},
            {"Field": "amount", "Offset": 11, "Length": 13, "Type": "DECIMAL", "Required": "YES", "InterfaceVersion": "rev 6",
             "Description": "Amount, zero padded, explicit decimal point. Currency implied USD."},
        ],
        "changelog": [
            {"Revision": "rev 6", "Date": "2025-11-03", "Change": "Baseline layout, 23 bytes.", "Author": "Core Ledger"},
            {"Revision": "rev 7", "Date": "2026-09-16", "Change": "customerId 10 -> 12; add currencyCode; amount 13 -> 10. 25 bytes.", "Author": "Core Ledger"},
        ],
        "note": "No machine-readable contract exists for this interface. Producers and consumers implement it by hand.",
    }
    _write_json(DOCS / "change-request-data.json", cr | {"_synthetic": True, "document": "change-request.pdf"})
    _write_json(DOCS / "release-notes-data.json", notes | {"_synthetic": True, "document": "release-notes.docx"})
    _write_json(DOCS / "interface-control-data.json", icd | {"_synthetic": True, "document": "interface-control.xlsx"})

    _change_request_pdf(cr)
    _release_notes_docx(notes)
    _icd_xlsx(icd)


def _change_request_pdf(cr: dict) -> None:
    from fpdf import FPDF

    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 9)
    pdf.set_text_color(120, 120, 120)
    pdf.cell(0, 6, "ACME RETAIL GROUP - IT SERVICE MANAGEMENT - SYNTHETIC DEMO DOCUMENT", new_x="LMARGIN", new_y="NEXT")
    pdf.set_text_color(0, 0, 0)
    pdf.set_font("Helvetica", "B", 18)
    pdf.cell(0, 12, f"Change Request {cr['id']}", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 11)
    pdf.multi_cell(0, 6, cr["title"], new_x="LMARGIN", new_y="NEXT")
    pdf.ln(3)
    rows = [("Release", cr["release"]), ("Requested by", cr["requested_by"]), ("Change type", cr["change_type"]),
            ("Risk", cr["risk"]), ("CAB", cr["cab"]), ("Window", cr["window"]), ("Status", cr["status"])]
    for key, value in rows:
        pdf.set_font("Helvetica", "B", 10)
        pdf.cell(40, 7, key)
        pdf.set_font("Helvetica", "", 10)
        pdf.cell(0, 7, value, new_x="LMARGIN", new_y="NEXT")
    pdf.ln(3)
    pdf.set_font("Helvetica", "B", 12)
    pdf.cell(0, 8, "Components", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 10)
    for c in cr["components"]:
        pdf.cell(0, 6, f"- {c['component']}: {c['from']} -> {c['to']}  ({c['target']})", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(2)
    pdf.set_font("Helvetica", "B", 12)
    pdf.cell(0, 8, "Dependency", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 10)
    pdf.multi_cell(0, 6, cr["dependency"], new_x="LMARGIN", new_y="NEXT")
    pdf.ln(2)
    pdf.set_font("Helvetica", "B", 12)
    pdf.cell(0, 8, "Implementation plan", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 10)
    for i, step in enumerate(cr["implementation"], 1):
        pdf.multi_cell(0, 6, f"{i}. {step}", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(2)
    for heading, key in (("Backout plan", "backout"), ("Verification", "verification")):
        pdf.set_font("Helvetica", "B", 12)
        pdf.cell(0, 8, heading, new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", "", 10)
        pdf.multi_cell(0, 6, cr[key], new_x="LMARGIN", new_y="NEXT")
        pdf.ln(1)
    pdf.output(str(DOCS / "change-request.pdf"))


def _release_notes_docx(notes: dict) -> None:
    from docx import Document

    doc = Document()
    doc.add_paragraph("Acme Retail Group - Release Management - synthetic demo document")
    doc.add_heading(f"Release notes {notes['release']}: {notes['title']}", level=1)
    doc.add_paragraph(notes["summary"])
    doc.add_heading("Changed services", level=2)
    table = doc.add_table(rows=1, cols=5)
    table.style = "Light Grid Accent 1"
    for cell, text in zip(table.rows[0].cells, ("Component", "From", "To", "Team", "Deployed by")):
        cell.text = text
    for s in notes["changed_services"]:
        row = table.add_row().cells
        for cell, key in zip(row, ("component", "from", "to", "team", "deployed_by")):
            cell.text = s[key]
    doc.add_heading("Deployment dependencies", level=2)
    for d in notes["deployment_dependencies"]:
        doc.add_paragraph(d, style="List Bullet")
    doc.add_heading("Expected rollout order", level=2)
    for step in notes["rollout_order"]:
        doc.add_paragraph(step, style="List Number")
    doc.add_heading("Migration requirements", level=2)
    for m in notes["migration_requirements"]:
        doc.add_paragraph(m, style="List Bullet")
    doc.add_heading("Compatibility notes", level=2)
    for n in notes["compatibility_notes"]:
        doc.add_paragraph(n, style="List Bullet")
    doc.add_heading("Validation", level=2)
    doc.add_paragraph(notes["validated_in"])
    doc.save(str(DOCS / "release-notes.docx"))


def _icd_xlsx(icd: dict) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    header_fill = PatternFill("solid", fgColor="1F2937")
    header_font = Font(bold=True, color="FFFFFF")
    for index, (sheet, key) in enumerate((("LEDGREC rev 7", "rev7"), ("LEDGREC rev 6", "rev6"), ("Change log", "changelog"))):
        ws = wb.active if index == 0 else wb.create_sheet()
        ws.title = sheet
        rows = icd[key]
        headers = list(rows[0].keys())
        ws.append(headers)
        for cell in ws[1]:
            cell.fill, cell.font = header_fill, header_font
        for row in rows:
            ws.append([row[h] for h in headers])
        for col, h in zip("ABCDEFG", headers):
            ws.column_dimensions[col].width = 60 if h in ("Description", "Change") else 16
    meta = wb.create_sheet("About")
    meta.append(["Document", icd["id"]])
    meta.append(["Title", icd["title"]])
    meta.append(["Owner", icd["owner"]])
    meta.append(["Note", icd["note"]])
    meta.append(["Synthetic", "All content is fictional demo data."])
    wb.save(str(DOCS / "interface-control.xlsx"))


# ================================================================================= main

def build(force: bool = False, quiet: bool = False) -> dict:
    commits = build_repos(force=force)
    build_adapters(commits)
    build_history(commits)
    build_documents(commits)
    _write_json(MERIDIAN / "repos" / "commits.json", commits)

    # Normalized snapshots, produced by the engine from the raw adapter outputs.
    from engine.reconciliation.reconciler import reconstruct_environment
    for env in S.ENVIRONMENTS:
        reality = reconstruct_environment(env)
        _write_json(ENVS / env / "snapshot.json", {
            "_generated": "Normalized by engine.reconciliation from environments/%s/adapters/*.json" % env,
            **reality.to_dict(),
        })
    if not quiet:
        print("Scenario built:")
        for comp, versions in commits.items():
            print(f"  {comp:16} " + "  ".join(f"{v}={h}" for v, h in versions.items()))
    return commits


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--force", action="store_true", help="rebuild git repositories")
    args = parser.parse_args()
    build(force=args.force)
