"""
Implicit contract discovery.

UNDERSTAND step. The LEDGER.IN interface has no machine-readable contract:
it exists only as a hand-maintained copybook on one side, layout constants on
the other, and an Interface Control Document spreadsheet. This module
reconstructs that implicit contract from the code at the exact deployed
commits plus the enterprise documents, and turns it into a probe spec.

Division of labour:
  * IBM Bob's contract-discovery agent (skill `implicit-contract-discovery`)
    performs the same investigation agentically and records its findings via
    the MCP tool `record_agent_finding`. Those findings are attached here.
  * This module is the deterministic extractor the product falls back on.
  * Neither of them decides compatibility. Only the executed probe does.
"""

from __future__ import annotations

import ast
import hashlib
import json
import re
from pathlib import Path

from engine.domain.models import DependencyEdge
from engine.paths import AGENT_FINDINGS, DOCUMENTS
from engine.probes import workspace
from engine.validation import memory

_PIC = re.compile(r"05\s+LR-([A-Z-]+)\s+PIC\s+X\((\d+)\)")


def _norm(name: str) -> str:
    return name.lower().replace("-", "_").replace("customerid", "customer_id").replace("currencycode", "currency_code")


def _line_of(text: str, needle: str) -> int | None:
    for number, line in enumerate(text.splitlines(), 1):
        if needle in line:
            return number
    return None


# ------------------------------------------------------------------ producer side

def producer_layouts(source: str, path: str, ref: str) -> dict:
    """Extract LAYOUT / LEGACY_LAYOUT tuples and version constants from the producer module."""
    tree = ast.parse(source)
    constants: dict[str, object] = {}
    layouts: dict[str, list[dict]] = {}
    for node in tree.body:
        if isinstance(node, ast.Assign) and len(node.targets) == 1 and isinstance(node.targets[0], ast.Name):
            name = node.targets[0].id
            try:
                value = ast.literal_eval(node.value)
            except ValueError:
                continue
            constants[name] = value
            if name.endswith("LAYOUT") and isinstance(node.value, ast.Tuple):
                fields, offset = [], 1
                for element in node.value.elts:
                    field, width = ast.literal_eval(element)
                    fields.append({
                        "name": _norm(field), "width": width, "offset": offset,
                        "source": f"{ref}:{path}#L{element.lineno}",
                    })
                    offset += width
                layouts[name] = fields
    out = {
        "primary": {
            "name": constants.get("LAYOUT_VERSION", "LAYOUT"),
            "fields": layouts.get("LAYOUT", []),
            "length": sum(f["width"] for f in layouts.get("LAYOUT", [])),
            "source": f"{ref}:{path}#L{_line_of(source, 'LAYOUT = (')}",
        },
        "compat": None,
        "hold_queue": constants.get("HOLD_QUEUE"),
    }
    if "LEGACY_LAYOUT" in layouts:
        out["compat"] = {
            "name": constants.get("LEGACY_LAYOUT_VERSION", "LEGACY_LAYOUT"),
            "fields": layouts["LEGACY_LAYOUT"],
            "length": sum(f["width"] for f in layouts["LEGACY_LAYOUT"]),
            "source": f"{ref}:{path}#L{_line_of(source, 'LEGACY_LAYOUT = (')}",
            "config_key": "ledger.recordLayout = v6-compat",
            "hold_rule": _line_of(source, "reasons.append") and f"{ref}:{path}#L{_line_of(source, 'reasons.append')}",
        }
    return out


# ------------------------------------------------------------------ consumer side

def copybook_fields(text: str, path: str, ref: str) -> dict:
    fields, offset = [], 1
    for number, line in enumerate(text.splitlines(), 1):
        match = _PIC.search(line)
        if match:
            width = int(match.group(2))
            fields.append({
                "name": _norm(match.group(1)), "width": width, "offset": offset,
                "pic": f"PIC X({width})", "source": f"{ref}:{path}#L{number}",
            })
            offset += width
    return {"name": Path(path).stem, "fields": fields, "length": sum(f["width"] for f in fields),
            "source": f"{ref}:{path}"}


def consumer_behaviour(source: str, path: str, ref: str) -> list[dict]:
    """Behaviour of the consumer that matters for compatibility, with line evidence."""
    facts = []
    if "_BY_LENGTH" in source or "len(record)" in source:
        line = _line_of(source, "_BY_LENGTH") or _line_of(source, "len(record)")
        facts.append({"id": "layout-selection", "value": "Selects the copybook by record length; unknown lengths are rejected.",
                      "source": f"{ref}:{path}#L{line}", "risk": "low"})
    else:
        line = _line_of(source, "record[offset:offset + width]") or _line_of(source, "def parse")
        facts.append({"id": "layout-selection",
                      "value": "Slices every record with one copybook. Record length is never checked; extra bytes are ignored.",
                      "source": f"{ref}:{path}#L{line}", "risk": "high"})
    match = re.search(r'DEFAULT_CURRENCY\s*=\s*"([A-Z]{3})"', source)
    if match:
        facts.append({"id": "default-currency",
                      "value": f"A record without a currency field is booked in {match.group(1)}.",
                      "source": f"{ref}:{path}#L{_line_of(source, 'DEFAULT_CURRENCY =')}", "risk": "medium"})
    if "re.sub(r\"[^0-9.]\"" in source or "[^0-9.]" in source:
        facts.append({"id": "tolerant-numeric",
                      "value": "Amounts are converted leniently: non-numeric bytes are stripped instead of rejected.",
                      "source": f"{ref}:{path}#L{_line_of(source, '[^0-9.]')}", "risk": "high"})
    if "\"status\": \"POSTED\"" in source:
        facts.append({"id": "success-signal",
                      "value": "Returns POSTED with MQ ACK and DB COMMIT for any record it can slice.",
                      "source": f"{ref}:{path}#L{_line_of(source, 'POSTED')}", "risk": "info"})
    return facts


# ------------------------------------------------------------------ documents

def icd_layouts() -> dict:
    """Read the Interface Control Document spreadsheet (LEDG-ICD-007)."""
    path = DOCUMENTS / "interface-control.xlsx"
    result = {"document": "documents/interface-control.xlsx", "sheets": {}}
    if not path.exists():
        result["status"] = "MISSING"
        return result
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    for sheet in wb.sheetnames:
        if not sheet.startswith("LEDGREC"):
            continue
        rows = list(wb[sheet].iter_rows(values_only=True))
        header = [str(h) for h in rows[0]]
        fields = []
        for index, row in enumerate(rows[1:], 2):
            record = dict(zip(header, row))
            fields.append({
                "name": _norm(str(record["Field"])), "offset": int(record["Offset"]), "width": int(record["Length"]),
                "type": record["Type"], "required": record["Required"], "description": record["Description"],
                "source": f"interface-control.xlsx!'{sheet}'!A{index}",
            })
        result["sheets"][sheet] = fields
    result["status"] = "READ"
    return result


def release_note_facts() -> list[dict]:
    path = DOCUMENTS / "release-notes.docx"
    if not path.exists():
        return []
    from docx import Document

    facts = []
    for index, paragraph in enumerate(Document(str(path)).paragraphs):
        text = paragraph.text.strip()
        if "legacy-ledger" in text and ("mq-bridge" in text or "rev" in text):
            facts.append({"text": text, "source": f"release-notes.docx ¶{index + 1}"})
    return facts


def change_request() -> dict | None:
    path = DOCUMENTS / "change-request-data.json"
    if not path.exists():
        return None
    cr = json.loads(path.read_text(encoding="utf-8"))
    return {"id": cr["id"], "window": cr["window"], "status": cr["status"], "dependency": cr["dependency"],
            "components": cr["components"], "source": "change-request.pdf (CR-4471)"}


# ------------------------------------------------------------------ comparison

def _effective_consumer_layout(consumer: dict, record_length: int) -> tuple[dict | None, str]:
    by_length = any(f["id"] == "layout-selection" and f["risk"] == "low" for f in consumer["behaviour"])
    copybooks = consumer["copybooks"]
    if by_length:
        match = next((c for c in copybooks if c["length"] == record_length), None)
        if match is None:
            return None, f"No copybook is {record_length} bytes long: the record is rejected (visible failure)."
        return match, f"Record length {record_length} selects {match['name']}."
    primary = next((c for c in copybooks if c["name"] == "LEDGREC"), copybooks[0])
    if primary["length"] == record_length:
        return primary, f"Record length matches {primary['name']}."
    return primary, (f"The receiver reads the message with {primary['name']} ({primary['length']} bytes) and ignores the "
                     f"last {record_length - primary['length']} byte(s) without raising an error.")


def compare(producer_layout: dict, consumer: dict) -> tuple[list[dict], str, str]:
    consumer_layout, selection_note = _effective_consumer_layout(consumer, producer_layout["length"])
    constraints = []
    if consumer_layout is None:
        return constraints, "INCOMPATIBLE", selection_note
    by_name = {f["name"]: f for f in consumer_layout["fields"]}
    default_currency = next((f for f in consumer["behaviour"] if f["id"] == "default-currency"), None)
    for pf in producer_layout["fields"]:
        cf = by_name.get(pf["name"])
        if cf is None:
            detail = "The receiver has no such field"
            if pf["name"] == "currency_code" and default_currency:
                detail += f"; {default_currency['value']}"
            constraints.append({
                "field": pf["name"], "producer": f"{pf['width']} bytes at {pf['offset']}", "consumer": "absent",
                "compatible": False, "risk": "Value is silently dropped.", "detail": detail,
                "producer_source": pf["source"], "consumer_source": default_currency["source"] if default_currency else consumer_layout["source"],
            })
            continue
        same = pf["offset"] == cf["offset"] and pf["width"] == cf["width"]
        if same:
            risk, detail = "None.", "Offsets and widths match."
        elif pf["offset"] == cf["offset"] and cf["width"] < pf["width"]:
            risk = "Value is silently cut off."
            detail = f"The receiver reads {cf['width']} of the {pf['width']} bytes the sender writes."
        else:
            risk = "Fields are out of line; the receiver reads bytes that belong to other fields."
            detail = (f"Sender writes bytes {pf['offset']}-{pf['offset'] + pf['width'] - 1}; "
                      f"receiver reads {cf['offset']}-{cf['offset'] + cf['width'] - 1}.")
        constraints.append({
            "field": pf["name"], "producer": f"{pf['width']} bytes at {pf['offset']}",
            "consumer": f"{cf['width']} bytes at {cf['offset']} ({cf.get('pic', '')})".strip(),
            "compatible": same, "risk": risk, "detail": detail,
            "producer_source": pf["source"], "consumer_source": cf["source"],
        })
    compatible = all(c["compatible"] for c in constraints)
    return constraints, "COMPATIBLE" if compatible else "INCOMPATIBLE", selection_note


# ------------------------------------------------------------------ entry point

def _agent_findings(edge_id: str, pv: str, cv: str) -> list[dict]:
    out = []
    if AGENT_FINDINGS.exists():
        for path in sorted(AGENT_FINDINGS.glob("*.json")):
            try:
                finding = json.loads(path.read_text(encoding="utf-8"))
            except json.JSONDecodeError:
                continue
            target = finding.get("target", {})
            if finding.get("kind") == "contract" and target.get("edge_id") == edge_id and \
                    target.get("producer_version") in (None, pv) and target.get("consumer_version") in (None, cv):
                out.append(finding)
    return out


def discover(edge: DependencyEdge, producer_version: str, producer_commit: str,
             consumer_version: str, consumer_commit: str, producer_config: dict | None = None) -> dict:
    contract_id = "contract-" + hashlib.sha1(
        f"{edge.id}|{producer_commit}|{consumer_commit}|{json.dumps(producer_config or {}, sort_keys=True)}".encode()
    ).hexdigest()[:10]
    base = {
        "contract_id": contract_id, "edge_id": edge.id, "interface": edge.interface,
        "formal_contract": edge.contract,
        "producer": {"component": edge.producer, "version": producer_version, "commit": producer_commit},
        "consumer": {"component": edge.consumer, "version": consumer_version, "commit": consumer_commit},
        "agent_findings": _agent_findings(edge.id, producer_version, consumer_version),
    }
    if edge.probe_type != "fixed-width-record":
        return base | {
            "status": "NEEDS_HUMAN",
            "method": "none",
            "summary": (f"No contract extractor exists for {edge.interface}. "
                        "Meridian will not guess compatibility for this boundary."),
            "constraints": [], "probe_spec": None,
        }

    p_path = edge.producer_entrypoint.split(":")[0]
    c_path = edge.consumer_entrypoint.split(":")[0]
    p_ref = f"{edge.producer}@{producer_commit}"
    c_ref = f"{edge.consumer}@{consumer_commit}"
    try:
        p_source = workspace.show(edge.producer, producer_commit, p_path)
        c_source = workspace.show(edge.consumer, consumer_commit, c_path)
        copybooks = [
            copybook_fields(workspace.show(edge.consumer, consumer_commit, path), path, c_ref)
            for path in workspace.list_files(edge.consumer, consumer_commit, "copybooks/")
        ]
    except workspace.SourceUnavailable as exc:
        return base | {"status": "INCONCLUSIVE", "method": "static", "summary": str(exc),
                       "constraints": [], "probe_spec": None}

    producer = producer_layouts(p_source, p_path, p_ref)
    consumer = {"copybooks": copybooks, "behaviour": consumer_behaviour(c_source, c_path, c_ref)}

    compat_mode = (producer_config or {}).get("ledger.recordLayout") == "v6-compat" and producer["compat"]
    layout = producer["compat"] if compat_mode else producer["primary"]
    constraints, prediction, selection_note = compare(layout, consumer)

    icd = icd_layouts()
    docs = {
        "interface_control": icd,
        "release_notes": release_note_facts(),
        "change_request": change_request(),
    }
    summary_bits = [selection_note]
    if compat_mode:
        summary_bits.append(
            f"The sender runs in compatibility mode ({producer['compat']['config_key']}): events that "
            f"{producer['compat']['name']} can't store are sent to {producer['hold_queue']} instead."
        )
    bad = [c["field"] for c in constraints if not c["compatible"]]
    if bad:
        summary_bits.append("Fields at risk: " + ", ".join(bad) + ".")

    spec = build_probe_spec(edge, producer_version, producer_commit, consumer_version, consumer_commit,
                            producer_config or {}, contract_id)
    return base | {
        "status": "DISCOVERED",
        "method": "static extraction from deployed commits + ICD spreadsheet + release notes + CR",
        "prediction": prediction,
        "prediction_note": "Static prediction only. The executable probe decides.",
        "summary": " ".join(summary_bits),
        "producer_layout": layout,
        "producer_layouts": producer,
        "consumer_copybooks": copybooks,
        "consumer_behaviour": consumer["behaviour"],
        "constraints": constraints,
        "documents": docs,
        "probe_spec": spec,
    }


def build_probe_spec(edge: DependencyEdge, producer_version: str, producer_commit: str,
                     consumer_version: str, consumer_commit: str, producer_config: dict,
                     contract_id: str | None = None, producer_label: str | None = None) -> dict:
    fixture_set = memory.fixtures()
    fixtures = [{
        "id": f["id"], "label": f["label"], "event": f["event"],
        "expect": {
            "customer_id": f["event"]["customerId"],
            "currency": f["event"]["currencyCode"],
            "amount": f"{float(f['event']['amount']):.2f}",
        },
    } for f in fixture_set]
    fingerprint = hashlib.sha1(json.dumps(
        [edge.id, producer_commit, consumer_commit, producer_config, fixtures], sort_keys=True
    ).encode()).hexdigest()[:8]
    return {
        "probe_id": f"probe-{edge.producer}-{producer_version}-{edge.consumer}-{consumer_version}-{fingerprint}",
        "edge_id": edge.id,
        "interface": edge.interface,
        "probe_type": edge.probe_type,
        "contract_id": contract_id,
        "producer": {
            "component": edge.producer, "version": producer_version, "label": producer_label or producer_version,
            "commit": producer_commit, "entrypoint": edge.producer_entrypoint, "config": producer_config,
        },
        "consumer": {
            "component": edge.consumer, "version": consumer_version, "commit": consumer_commit,
            "entrypoint": edge.consumer_entrypoint,
        },
        "fixture_source": "Stage validation corpus (environments/validation-runs.json)",
        "fixtures": fixtures,
        "assertions": [
            "Every record the producer delivers is posted with exactly the source customer_id.",
            "Every delivered record is posted with exactly the source currency.",
            "Every delivered record is posted with exactly the source amount.",
            "A record the producer cannot represent must be held explicitly, never truncated.",
        ],
        "isolation": "git archive of each exact commit into a temporary sandbox; python -I subprocess; "
                     "no environment credentials; no network endpoints configured",
    }
